from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document

import project


def class_file_counts(root: Path) -> list[tuple[int, int]]:
    rows: list[tuple[int, int]] = []
    if not root.exists():
        return rows
    for folder in sorted([p for p in root.iterdir() if p.is_dir()], key=lambda p: int(p.name)):
        rows.append((int(folder.name), len(list(folder.glob("*.csv")))))
    return rows


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_paragraph(doc: Document, text: str) -> None:
    doc.add_paragraph(text)


def add_list(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def build_document() -> Document:
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    raw_root = project.RAW_DATA_DIR
    total_csv = len(list(raw_root.rglob("*.csv"))) if raw_root.exists() else 0
    class_counts = class_file_counts(raw_root)

    doc = Document()
    add_heading(doc, "AI Oil&Gas 3W Project Guide", level=0)
    add_paragraph(doc, f"Generated: {now}")

    add_heading(doc, "1. What We Built (Step by Step)")
    add_list(
        doc,
        [
            "Created project structure (data, models, reports, incoming).",
            "Loaded 3W dataset into data/raw/3w.",
            "Implemented full ML pipeline in project.py (feature extraction, training, evaluation, saving artifacts).",
            "Added single-file inference mode: python project.py --predict <file.csv>.",
            "Added human-readable class names/descriptions and text summary report.",
            "Built Streamlit mini app (app.py) for browser-based CSV upload and prediction.",
        ],
    )

    add_heading(doc, "2. Project Folder Structure")
    add_list(
        doc,
        [
            r"C:\AI\ai_oil_3w\project.py - train/predict pipeline",
            r"C:\AI\ai_oil_3w\app.py - mini web app",
            r"C:\AI\ai_oil_3w\data\raw\3w\0..8 - dataset by class",
            r"C:\AI\ai_oil_3w\incoming - uploaded/test CSV files",
            r"C:\AI\ai_oil_3w\models\rf_3w.joblib - trained model",
            r"C:\AI\ai_oil_3w\models\feature_columns.json - feature schema",
            r"C:\AI\ai_oil_3w\reports\* - metrics and prediction summaries",
        ],
    )

    add_heading(doc, "3. Dataset Summary (3W)")
    add_paragraph(doc, f"Total CSV files found: {total_csv}")
    add_paragraph(doc, "Core columns in each time-series CSV:")
    add_list(
        doc,
        [
            "timestamp",
            "P-PDG",
            "P-TPT",
            "T-TPT",
            "P-MON-CKP",
            "T-JUS-CKP",
            "P-JUS-CKGL",
            "T-JUS-CKGL",
            "QGL",
            "class (in labeled training files)",
        ],
    )

    add_heading(doc, "4. Class Definitions")
    class_table = doc.add_table(rows=1, cols=4)
    class_table.style = "Light List Accent 1"
    header_cells = class_table.rows[0].cells
    header_cells[0].text = "class_id"
    header_cells[1].text = "class_name"
    header_cells[2].text = "description"
    header_cells[3].text = "file_count"

    count_map = {cls: cnt for cls, cnt in class_counts}
    for cls in sorted(project.CLASS_NAME_MAP.keys()):
        row = class_table.add_row().cells
        row[0].text = str(cls)
        row[1].text = project.class_name(cls)
        row[2].text = project.class_description(cls)
        row[3].text = str(count_map.get(cls, 0))

    add_heading(doc, "5. Model and Features")
    add_paragraph(doc, "Model: RandomForestClassifier (scikit-learn).")
    add_paragraph(doc, "Features extracted from each sensor time-series:")
    add_list(
        doc,
        [
            "mean, std, min, max, median, p10, p90",
            "trend (linear slope)",
            "nan_ratio",
            "duration_sec and valid_timestamp_ratio",
        ],
    )

    add_heading(doc, "6. How to Run After PC Startup")
    add_paragraph(doc, "Recommended (browser app):")
    add_list(
        doc,
        [
            r"cd C:\AI\ai_oil_3w",
            r".\.venv\Scripts\Activate.ps1",
            "streamlit run app.py",
            "Open http://localhost:8501 in browser, upload CSV, click prediction button.",
        ],
    )
    add_paragraph(doc, "Command-line mode:")
    add_list(
        doc,
        [
            r"cd C:\AI\ai_oil_3w",
            r".\.venv\Scripts\Activate.ps1",
            r"python project.py --predict .\incoming\my_case.csv",
        ],
    )

    add_heading(doc, "7. Retraining")
    add_list(
        doc,
        [
            r"cd C:\AI\ai_oil_3w",
            r".\.venv\Scripts\Activate.ps1",
            "python project.py",
        ],
    )

    add_heading(doc, "8. Using Real (Non-Training) Data")
    add_list(
        doc,
        [
            "Prepare CSV with same sensor columns as training data.",
            "Put file into incoming folder.",
            "Run prediction from app.py or project.py --predict.",
            "Review reports/prediction_summary_*.txt and confidence level.",
        ],
    )

    add_heading(doc, "9. Important Limitations")
    add_list(
        doc,
        [
            "This is a decision-support model, not autonomous well control.",
            "Performance depends on similarity between real data and 3W training data.",
            "If site conditions differ, model adaptation/retraining is required.",
        ],
    )

    return doc


def main() -> None:
    project.REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_path = project.REPORTS_DIR / f"Project_Guide_{stamp}.docx"

    doc = build_document()
    doc.save(out_path)
    print(f"Word guide created: {out_path}")


if __name__ == "__main__":
    main()
